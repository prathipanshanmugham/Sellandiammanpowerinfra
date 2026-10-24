"""Generic record endpoints for every entity in the registry, plus files, approvals, search and exports."""
import csv
import io
import re
from fastapi import APIRouter, Request, Depends, HTTPException, UploadFile, File, Form, Query
from fastapi.responses import Response, StreamingResponse

import registry as REG
import engine as ENG
from catalog import OPTIONS, label_of
from core import (db, current_user, staff_user, can, get_settings, save_upload, sign_file, check_file_token, UPLOAD_DIR,
                  audit, now_iso, new_id, rate_limit, DOC_TYPES, IMAGE_TYPES, verify_audit_chain, require)
from roles import HIDDEN_FIELD_TAGS, PORTAL_ROLES, MODULES, ROLES
from pdf import Letterhead, ind, rs, words_inr, P

router = APIRouter()


# ------------------------------------------------------------------ meta
@router.get("/meta")
async def meta(user=Depends(current_user)):
    from roles import PERMS
    tags = HIDDEN_FIELD_TAGS.get(user["role"], [])
    perms = PERMS.get(user["role"], {})
    if user["role"] in PORTAL_ROLES:
        ents = {}
        for k, e in REG.E.items():
            if user["role"] in (e.get("portal") or {}):
                d = dict(e)
                d["fields"] = [ENG_strip(f, tags) for f in e["fields"] if f.get("hide") not in tags]
                d["perm"] = "r"
                ents[k] = d
    else:
        ents = REG.meta_for(user["role"], tags, perms)
    return {"entities": ents, "options": OPTIONS, "nav_order": REG.NAV_ORDER, "modules": MODULES,
            "roles": [{"value": k, "label": v} for k, v in ROLES.items()]}


def ENG_strip(f, tags):
    return REG._strip(f, tags)


# ------------------------------------------------------------------ records
@router.get("/e/{key}")
async def list_(key: str, request: Request, user=Depends(current_user)):
    params = dict(request.query_params)
    return await ENG.list_records(key, user, params)


@router.get("/e/{key}/{rid}")
async def get_(key: str, rid: str, user=Depends(current_user)):
    doc = await ENG.get_record(key, user, rid)
    e = ENG.ent(key)
    doc["_actions"] = ENG.available_actions(e, doc, user) if user["role"] not in PORTAL_ROLES else []
    if doc.get("approval_id"):
        doc["_approval"] = await db.approvals.find_one({"id": doc["approval_id"]}, {"_id": 0})
    return doc


@router.post("/e/{key}")
async def create_(key: str, body: dict, user=Depends(staff_user)):
    return await ENG.create_record(key, user, body)


@router.patch("/e/{key}/{rid}")
async def update_(key: str, rid: str, body: dict, user=Depends(staff_user)):
    return await ENG.update_record(key, user, rid, body)


@router.delete("/e/{key}/{rid}")
async def delete_(key: str, rid: str, reason: str = "", user=Depends(staff_user)):
    return await ENG.delete_record(key, user, rid, reason)


@router.post("/e/{key}/{rid}/action/{action}")
async def action_(key: str, rid: str, action: str, body: dict | None = None, user=Depends(staff_user)):
    out = await ENG.run_action(key, user, rid, action, body or {})
    out["_actions"] = ENG.available_actions(ENG.ent(key), out, user)
    return out


@router.get("/e/{key}/{rid}/history")
async def history_(key: str, rid: str, user=Depends(staff_user)):
    return await ENG.history(key, user, rid)


@router.post("/e/{key}/bulk")
async def bulk_(key: str, body: dict, user=Depends(staff_user)):
    """Bulk change of one field (status for simple records, owner, etc.) or bulk delete."""
    ids = (body.get("ids") or [])[:500]
    done, errors = 0, []
    for rid in ids:
        try:
            if body.get("op") == "delete":
                await ENG.delete_record(key, user, rid, body.get("reason") or "")
            elif body.get("op") == "set":
                await ENG.update_record(key, user, rid, body.get("values") or {})
            done += 1
        except HTTPException as ex:
            errors.append({"id": rid, "error": ex.detail})
    return {"done": done, "errors": errors[:20]}


# ------------------------------------------------------------------ exports
def cell_text(e, f, v, refs):
    if v is None:
        return ""
    t = f["t"] if f else "text"
    if t in ("ref",):
        return (refs.get(f["ref"]) or {}).get(v, "")
    if t == "refs":
        return ", ".join((refs.get(f["ref"]) or {}).get(x, "") for x in v or [])
    if t == "select":
        o = ENG.opts_of(f)
        return next((x["label"] for x in o if x["value"] == v), v)
    if t == "multiselect":
        o = ENG.opts_of(f)
        return ", ".join(next((x["label"] for x in o if x["value"] == y), y) for y in v or [])
    if t == "bool":
        return "Yes" if v else "No"
    if t in ("money",):
        return ind(v)
    if t in ("lines", "json", "files", "file", "geo"):
        return f"{len(v)} lines" if isinstance(v, list) else ""
    return str(v)


@router.get("/e/{key}/export/{fmt}")
async def export_(key: str, fmt: str, request: Request, user=Depends(staff_user)):
    e = ENG.ent(key)
    if not can(user, e["module"], "e") and not can(user, "reports", "e"):
        raise HTTPException(403, "You don't have permission to export this")
    params = dict(request.query_params)
    params["size"] = "5000"
    data = await ENG.list_records(key, user, params)
    fmap = {f["k"]: f for f in e["fields"]}
    cols = [c for c in (params.get("cols", "").split(",") if params.get("cols") else e["list"]) if c in fmap or c in ("code", "created_at", "status")]
    sf = ENG.status_field(e)
    heads = [("Code" if c == "code" else "Created" if c == "created_at" else "Status" if c == "status" else fmap[c]["l"]) for c in cols]
    rows = []
    for d in data["items"]:
        r = []
        for c in cols:
            if c == "status" and sf:
                st = next((s["label"] for s in e["status"]["states"] if s["value"] == d.get(sf)), d.get(sf))
                r.append(st or "")
            elif c == "created_at":
                r.append((d.get("created_at") or "")[:10])
            else:
                r.append(cell_text(e, fmap.get(c), d.get(c), data["refs"]))
        rows.append(r)
    await audit(user, "export", key, None, {"format": fmt, "rows": len(rows)})
    name = f"{key}-{now_iso()[:10]}"
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(heads)
        w.writerows(rows)
        return Response(buf.getvalue().encode("utf-8-sig"), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{name}.csv"'})
    if fmt == "xlsx":
        return Response(xlsx(heads, rows, e["plural"]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})
    if fmt == "pdf":
        s = await get_settings()
        doc = Letterhead(e["plural"], s, subtitle=f"{len(rows)} records, exported {now_iso()[:10]} by {user.get('name')}", landscape_mode=len(cols) > 6)
        align = ["R" if (fmap.get(c) or {}).get("t") in ("money", "number", "percent") else "L" for c in cols]
        doc.table(heads, rows, align=align)
        return Response(doc.build(), media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{name}.pdf"'})
    raise HTTPException(404, "Unknown format")


def xlsx(heads, rows, title="Sheet"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[^\w ]", "", title)[:30] or "Sheet"
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0E1726")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        out = []
        for v in r:
            if isinstance(v, str) and re.fullmatch(r"-?[\d,]+(\.\d+)?", v) and "," in v:
                try:
                    out.append(float(v.replace(",", "")))
                    continue
                except ValueError:
                    pass
            out.append(v)
        ws.append(out)
    for i, h in enumerate(heads, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200]]) + 2
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(48, width)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ record PDFs (invoice, PO, proposal, RA bill, test certificate...)
@router.get("/e/{key}/{rid}/pdf")
async def record_pdf(key: str, rid: str, user=Depends(current_user)):
    doc = await ENG.get_record(key, user, rid)
    s = await get_settings()
    pdf = await build_record_pdf(key, doc, s, user)
    await audit(user, "pdf", key, rid, {"code": doc.get("code")})
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{(doc.get("code") or key).replace("/", "-")}.pdf"'})


async def ref_name(coll, rid, field="name"):
    if not rid:
        return ""
    r = await db[coll].find_one({"id": rid}, {"_id": 0})
    if not r:
        return ""
    return r.get(field) or r.get("name") or r.get("title") or r.get("code") or ""


TITLES = {"invoices": "Tax invoice", "purchase_orders": "Purchase order", "proposals": "Techno-commercial proposal", "ra_bills": "Running account bill",
          "test_reports": "Test certificate", "permits": "Permit to work", "work_orders": "Work order", "receipts": "Receipt", "boqs": "Estimate",
          "pm_jobs": "Maintenance report", "tickets": "Service report", "grns": "Goods receipt note", "material_moves": "Material movement",
          "variations": "Variation / extra items", "sc_bills": "Subcontractor bill", "rfqs": "Request for quotation", "dprs": "Daily progress report",
          "handovers": "Handover certificate", "vendor_payments": "Payment voucher", "credit_notes": "Credit note", "inspections": "Inspection report"}


async def build_record_pdf(key, doc, s, user):
    e = ENG.ent(key)
    refs = doc.get("_refs") or {}
    title = TITLES.get(key, e["label"])
    if key == "invoices" and doc.get("itype") == "proforma":
        title = "Proforma invoice"
    sf = ENG.status_field(e)
    watermark = None
    if sf and doc.get(sf) in ("draft", "pending_approval"):
        watermark = "DRAFT"
    if sf and doc.get(sf) in ("cancelled",):
        watermark = "CANCELLED"
    L = Letterhead(title, s, subtitle=doc.get("code") or ENG.title_of(e, doc), watermark=watermark)
    fmap = {f["k"]: f for f in e["fields"]}
    tags = HIDDEN_FIELD_TAGS.get(user["role"], [])
    if key == "invoices":
        acc = await db.accounts.find_one({"id": doc.get("account_id")}, {"_id": 0}) or {}
        L.meta([("Invoice no.", doc.get("code")), ("Date", doc.get("date")), ("Bill to", acc.get("name")), ("Due date", doc.get("due_date")),
                ("Client GSTIN", acc.get("gstin")), ("Place of supply", doc.get("place_of_supply")), ("Address", acc.get("address")),
                ("Project", refs.get("projects", {}).get(doc.get("project_id"))), ("IRN", doc.get("irn")), ("E-way bill", doc.get("eway_bill"))])
        rows = [[n, ln.get("description"), ln.get("hsn"), ind(ln.get("qty"), 2), ln.get("unit"), ind(ln.get("rate")), f"{ln.get('gst_rate')}%", ind(ln.get("amount"))]
                for n, ln in enumerate(doc.get("lines") or [], 1)]
        L.table(["#", "Description", "HSN/SAC", "Qty", "Unit", "Rate", "GST", "Amount"], rows, widths=[4, 40, 9, 8, 6, 11, 6, 13], align=list("LLLRLRRR"))
        tot = [("Taxable value", doc.get("subtotal"))]
        if doc.get("interstate"):
            tot.append(("IGST", doc.get("igst")))
        else:
            tot += [("CGST", doc.get("cgst")), ("SGST", doc.get("sgst"))]
        tot.append(("Total", doc.get("total")))
        if doc.get("retention"):
            tot.append(("Less retention", doc.get("retention")))
        if doc.get("tds_expected"):
            tot.append(("Less TDS (expected)", doc.get("tds_expected")))
        if doc.get("retention") or doc.get("tds_expected"):
            tot.append(("Amount due", doc.get("receivable")))
        L.totals(tot, words_inr(doc.get("total")))
        if s.get("bank_details"):
            L.h("Bank details")
            L.p(s["bank_details"])
        L.signature("Received by (client)")
        return L.build()
    if key == "proposals":
        acc = await db.accounts.find_one({"id": doc.get("account_id")}, {"_id": 0}) or {}
        L.meta([("Proposal no.", doc.get("code")), ("Date", (doc.get("sent_at") or doc.get("created_at") or "")[:10]), ("To", acc.get("name")),
                ("Valid until", doc.get("valid_until")), ("Subject", doc.get("title"))])
        for k, h in (("cover_letter", None), ("scope", "Scope of work")):
            if doc.get(k):
                if h:
                    L.h(h)
                L.md(doc[k])
        if doc.get("boq_id"):
            b = await db.boqs.find_one({"id": doc["boq_id"]}, {"_id": 0}) or {}
            rows = [[n, ln.get("description"), ln.get("make") or "", ind(ln.get("qty"), 2), ln.get("unit"), ind(ln.get("rate")), ind(ln.get("amount"))]
                    for n, ln in enumerate(b.get("lines") or [], 1)]
            L.h("Bill of quantities")
            L.table(["#", "Description", "Make", "Qty", "Unit", "Rate", "Amount"], rows, widths=[4, 44, 10, 8, 6, 12, 14], align=list("LLLRLRR"))
        tot = [("Subtotal", doc.get("subtotal"))]
        if doc.get("discount_pct"):
            tot.append((f"Discount {doc['discount_pct']:g}%", round(float(doc.get("subtotal") or 0) - float(doc.get("taxable") or 0), 2)))
        tot += [("GST", doc.get("gst")), ("Total", doc.get("total"))]
        L.totals(tot, words_inr(doc.get("total")))
        for k, h in (("exclusions", "Exclusions"), ("deviations", "Deviations"), ("payment_terms", "Payment terms"), ("compliance", "Compliance")):
            if doc.get(k):
                L.h(h)
                L.md(doc[k])
        terms = await db.templates.find_one({"kind": "terms", "name": {"$regex": doc.get("terms_set") or "standard", "$options": "i"}, "deleted": {"$ne": True}}, {"_id": 0})
        if terms:
            L.h("Terms & conditions")
            L.md(terms.get("body"))
        L.signature("Accepted by (client)")
        return L.build()
    if key == "test_reports":
        ins = await db.instruments.find_one({"id": doc.get("instrument_id")}, {"_id": 0}) or {}
        L.meta([("Certificate no.", doc.get("code")), ("Date of test", doc.get("date")), ("Test", label_of("test_types", doc.get("test_type"))),
                ("Equipment", doc.get("equipment")), ("Project", refs.get("projects", {}).get(doc.get("project_id"))),
                ("Instrument", f"{ins.get('name', '')}, S/N {ins.get('serial_no', '')}, calibrated until {ins.get('calibration_due', '')}"),
                ("Result", (doc.get("result") or "").upper()), ("Witnessed by", doc.get("witnessed_by"))])
        rows = [[ln.get("point"), ln.get("value"), ln.get("unit"), ln.get("limit"), (ln.get("result") or "").upper()] for ln in doc.get("readings") or []]
        L.table(["Point / phase", "Measured", "Unit", "Acceptance limit", "Result"], rows, widths=[30, 18, 10, 26, 12])
        L.p(doc.get("remarks"))
        L.signature("Tested by: " + (refs.get("users", {}).get(doc.get("tested_by")) or ""), "Witness")
        return L.build()
    if key == "permits":
        L.meta([("Permit no.", doc.get("code")), ("Type", label_of("permit_types", doc.get("ptype"))), ("Project", refs.get("projects", {}).get(doc.get("project_id"))),
                ("Location", doc.get("location")), ("Equipment / circuit", doc.get("equipment")), ("Valid", f"{doc.get('valid_from', '')} → {doc.get('valid_to', '')}"),
                ("Authorised person", refs.get("users", {}).get(doc.get("authorised_person_id"))), ("Status", doc.get("status"))])
        L.h("Work")
        L.p(doc.get("description"))
        L.h("Precautions")
        L.table(["Check", "Done"], [[ln.get("check"), "Yes" if ln.get("done") else "No"] for ln in doc.get("precautions") or []], widths=[80, 20])
        if doc.get("isolations"):
            L.h("Isolation points (LOTO)")
            L.table(["Point", "Lock no.", "Tag no.", "Locked by"], [[x.get("point"), x.get("lock_no"), x.get("tag_no"), x.get("locked_by")] for x in doc["isolations"]])
        if doc.get("closure_note"):
            L.h("Closure")
            L.p(doc["closure_note"])
        L.signature("Permit receiver", "Permit issuer")
        return L.build()
    # generic: header fields + every lines table + money totals
    pairs = []
    for f in e["fields"]:
        if f["t"] in ("lines", "markdown", "files", "file", "json", "geo", "textarea") or f.get("hide") in tags:
            continue
        v = doc.get(f["k"])
        if v in (None, "", [], False):
            continue
        if f["t"] == "money" and f.get("calc"):
            continue
        txt = cell_text(e, f, v, refs) if f["t"] != "money" else rs(v)
        pairs.append((f["l"], txt))
    if sf:
        pairs.insert(0, ("Status", next((x["label"] for x in e["status"]["states"] if x["value"] == doc.get(sf)), doc.get(sf))))
    L.meta(pairs[:24])
    for f in e["fields"]:
        if f["t"] in ("textarea", "markdown") and doc.get(f["k"]) and f.get("hide") not in tags:
            L.h(f["l"])
            L.md(doc[f["k"]])
    for f in e["fields"]:
        if f["t"] != "lines" or not doc.get(f["k"]) or f.get("hide") in tags:
            continue
        sub = [x for x in f["fields"] if x.get("hide") not in tags and x["t"] not in ("files", "file")]
        heads = [x["l"] for x in sub]
        rows = [[cell_text(e, x, ln.get(x["k"]), refs) if x["t"] not in ("number",) else ind(ln.get(x["k"]), 2).rstrip("0").rstrip(".") if ln.get(x["k"]) not in (None, "") else "" for x in sub]
                for ln in doc[f["k"]]]
        L.h(f["l"])
        L.table(heads, rows, align=["R" if x["t"] in ("money", "number", "percent") else "L" for x in sub])
    money_calc = [(f["l"], doc.get(f["k"])) for f in e["fields"] if f["t"] == "money" and f.get("calc") and f.get("hide") not in tags and doc.get(f["k"]) not in (None, 0)]
    if money_calc:
        L.totals(money_calc[-6:])
    L.signature()
    return L.build()


# ------------------------------------------------------------------ files
@router.post("/files")
async def upload(file: UploadFile = File(...), public: bool = Form(False), user=Depends(staff_user)):
    rec = await save_upload(file, public=public)
    return {"id": rec["id"], "name": rec["original_name"], "type": rec["content_type"], "size": rec["size"],
            "url": f"/api/public/media/{rec['id']}" if public else None}


@router.post("/portal/files")
async def portal_upload(request: Request, file: UploadFile = File(...), user=Depends(current_user)):
    rate_limit(request, "pupload", 30, 600)
    rec = await save_upload(file, max_bytes=15 * 1024 * 1024)
    return {"id": rec["id"], "name": rec["original_name"], "type": rec["content_type"], "size": rec["size"]}


@router.get("/files/{fid}/link")
async def file_link(fid: str, user=Depends(current_user)):
    f = await db.files.find_one({"id": fid}, {"_id": 0})
    if not f:
        raise HTTPException(404, "File not found")
    return {"url": f"/api/files/get/{sign_file(fid)}"}


@router.get("/files/get/{token}")
async def file_get(token: str):
    fid = check_file_token(token)
    f = await db.files.find_one({"id": fid}, {"_id": 0})
    p = UPLOAD_DIR / fid
    if not f or not p.exists():
        raise HTTPException(404, "File not found")
    return Response(p.read_bytes(), media_type=f["content_type"],
                    headers={"Content-Disposition": f'inline; filename="{f["original_name"]}"', "Cache-Control": "private, max-age=600"})


@router.get("/public/media/{fid}")
async def media(fid: str):
    f = await db.files.find_one({"id": fid, "public": True}, {"_id": 0})
    p = UPLOAD_DIR / fid
    if not f or not p.exists():
        raise HTTPException(404, "Not found")
    return Response(p.read_bytes(), media_type=f["content_type"], headers={"Cache-Control": "public, max-age=31536000, immutable"})


# ------------------------------------------------------------------ approvals inbox
@router.get("/approvals/inbox")
async def approvals_inbox(user=Depends(staff_user), status: str = "pending"):
    rows = await db.approvals.find({"status": status} if status != "all" else {}, {"_id": 0}).sort("created_at", -1).to_list(500)
    mine, others = [], []
    for r in rows:
        ok, why = await ENG.can_decide(user, r)
        r["can_decide"] = ok
        r["why_not"] = why
        r["entity_label"] = REG.E.get(r["doc_entity"], {}).get("label")
        (mine if ok else others).append(r)
    asked = [r for r in rows if r.get("requested_by") == user["id"]]
    names = {u["id"]: u["name"] for u in await db.users.find({"id": {"$in": list({r["requested_by"] for r in rows})}}, {"_id": 0, "id": 1, "name": 1}).to_list(500)}
    for r in rows:
        r["requested_by_name"] = names.get(r["requested_by"])
    return {"to_decide": mine, "my_requests": asked, "count": len(mine)}


@router.post("/approvals/{aid}/decide")
async def approvals_decide(aid: str, body: dict, user=Depends(staff_user)):
    return await ENG.decide(user, aid, body.get("decision"), body.get("reason") or "")


# ------------------------------------------------------------------ global search (Ctrl+K)
SEARCHABLE = ["leads", "accounts", "contacts", "projects", "proposals", "boqs", "tenders", "purchase_orders", "invoices", "vendors", "items",
              "tickets", "installed_assets", "employees", "documents", "liaisons", "drawings", "cable_drums", "serial_items", "workers"]


@router.get("/search")
async def search(q: str, user=Depends(staff_user)):
    q = (q or "").strip()
    if len(q) < 2:
        return []
    out = []
    for key in SEARCHABLE:
        e = REG.E[key]
        if not can(user, e["module"], "r"):
            continue
        try:
            res = await ENG.list_records(key, user, {"q": q, "size": "5"})
        except HTTPException:
            continue
        for d in res["items"]:
            out.append({"entity": key, "entity_label": e["label"], "id": d["id"], "title": ENG.title_of(e, d), "code": d.get("code"),
                        "icon": e.get("icon")})
    return out[:40]


# ------------------------------------------------------------------ saved views
@router.get("/views/{key}")
async def views(key: str, user=Depends(staff_user)):
    return await db.saved_views.find({"entity": key, "$or": [{"user_id": user["id"]}, {"shared": True}]}, {"_id": 0}).to_list(50)


@router.post("/views/{key}")
async def save_view(key: str, body: dict, user=Depends(staff_user)):
    doc = {"id": new_id(), "entity": key, "user_id": user["id"], "name": str(body.get("name") or "View")[:40],
           "params": {k: str(v)[:100] for k, v in (body.get("params") or {}).items()}, "cols": (body.get("cols") or [])[:20],
           "shared": bool(body.get("shared")), "created_at": now_iso()}
    await db.saved_views.insert_one(dict(doc))
    return doc


@router.delete("/views/{key}/{vid}")
async def delete_view(key: str, vid: str, user=Depends(staff_user)):
    await db.saved_views.delete_one({"id": vid, "user_id": user["id"]})
    return {"ok": True}


# ------------------------------------------------------------------ activity log
@router.get("/audit")
async def audit_log(user=Depends(require("audit", "r")), entity: str = "", q: str = "", page: int = 1):
    f = {}
    if entity:
        f["entity"] = entity
    if q:
        rx = {"$regex": re.escape(q[:40]), "$options": "i"}
        f["$or"] = [{"user_name": rx}, {"action": rx}, {"details.code": rx}]
    total = await db.audit_logs.count_documents(f)
    rows = await db.audit_logs.find(f, {"_id": 0}).sort("seq", -1).skip((page - 1) * 100).limit(100).to_list(100)
    return {"items": rows, "total": total}


@router.get("/audit/verify")
async def audit_verify(user=Depends(require("audit", "r"))):
    return await verify_audit_chain()
